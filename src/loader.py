"""Bounded, non-extracting upload readers. Content errors are collected later."""
from io import BytesIO, StringIO
from pathlib import PurePosixPath
from zipfile import ZipFile
import csv

import pandas as pd

SCHEMAS = {
    "crm_orders": ["order_id", "client_name", "dealer", "manager", "order_date",
                   "planned_delivery", "sale_amount", "crm_status"],
    "measurements": ["order_id", "measurer", "planned_date", "actual_date",
                     "measurement_status", "design_approved", "specification_ready"],
    "production": ["order_id", "production_stage", "production_start", "planned_finish",
                   "actual_finish", "materials_cost", "labor_cost", "other_cost"],
    "installation": ["order_id", "installer", "planned_date", "actual_date",
                     "installation_status", "installation_cost", "delivery_cost", "complaint"],
}
FILENAMES = {source: source + (".csv" if source == "crm_orders" else ".xlsx")
             for source in SCHEMAS}
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ZIP_BYTES = 20 * 1024 * 1024
MAX_TOTAL_BYTES = 40 * 1024 * 1024
MAX_XLSX_EXPANDED = 50 * 1024 * 1024
MAX_ROWS = 50_000
MAX_ARCHIVE_MEMBERS = 100
MAX_COLUMNS = 100


class LoadError(ValueError):
    """An actionable input error safe to show in the interface."""


def _safe_members(archive, expanded_limit):
    members = archive.infolist()
    if len(members) > MAX_ARCHIVE_MEMBERS:
        raise LoadError("Слишком много файлов внутри архива (максимум 100).")
    if sum(info.file_size for info in members) > expanded_limit:
        raise LoadError("Превышен допустимый распакованный размер архива.")
    for info in members:
        path = PurePosixPath(info.filename)
        if (path.is_absolute() or ".." in path.parts or "\\" in info.filename
                or ":" in info.filename or "\x00" in info.filename):
            raise LoadError("Архив содержит небезопасный путь.")
        if ((info.external_attr >> 16) & 0o170000) == 0o120000:
            raise LoadError("Символические ссылки внутри архива запрещены.")
        if info.flag_bits & 1:
            raise LoadError("Зашифрованные архивы не поддерживаются.")
    return members


def read_source(source: str, data: bytes) -> pd.DataFrame:
    if source not in SCHEMAS:
        raise LoadError("Неизвестный источник данных.")
    filename = FILENAMES[source]
    if not data:
        raise LoadError(f"{filename}: файл пуст.")
    if len(data) > MAX_FILE_BYTES:
        raise LoadError(f"{filename}: размер превышает 10 МБ.")
    try:
        if source == "crm_orders":
            # Preserve IDs (including leading zeros), blanks and unparsed values.
            decoded = data.decode("utf-8-sig")
            delimiter = ";" if decoded.splitlines()[0].count(";") > decoded.splitlines()[0].count(",") else ","
            records = csv.reader(StringIO(decoded), delimiter=delimiter, strict=True)
            header = next(records)
            if len(header) > MAX_COLUMNS:
                raise LoadError(f"{filename}: более {MAX_COLUMNS} столбцов.")
            values = []
            for row in records:
                if not row:
                    continue
                if len(row) != len(header):
                    raise LoadError(f"{filename}: количество полей не совпадает с заголовком в записи {len(values) + 2}.")
                values.append(row)
                if len(values) > MAX_ROWS:
                    raise LoadError(f"{filename}: более {MAX_ROWS} строк.")
            frame = pd.DataFrame(values, columns=header)
        else:
            with ZipFile(BytesIO(data)) as xlsx:
                _safe_members(xlsx, MAX_XLSX_EXPANDED)
            from openpyxl import load_workbook
            workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
            try:
                sheet = workbook.worksheets[0]
                if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                    raise LoadError(f"{filename}: более {MAX_COLUMNS} столбцов.")
                if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
                    raise LoadError(f"{filename}: более {MAX_ROWS} строк.")
                rows = sheet.iter_rows(values_only=True)
                header = list(next(rows, ()))
                values = []
                for row in rows:
                    if len(values) >= MAX_ROWS:
                        raise LoadError(f"{filename}: более {MAX_ROWS} строк.")
                    values.append(row)
                frame = pd.DataFrame(values, columns=header)
            finally:
                workbook.close()
        if len(frame) > MAX_ROWS:
            raise LoadError(f"{filename}: более {MAX_ROWS} строк.")
        normalized_header = [str(column).strip() for column in header]
        if len(normalized_header) != len(set(normalized_header)):
            raise LoadError(f"{filename}: повторяющиеся названия столбцов.")
        frame.columns = normalized_header
        missing = set(SCHEMAS[source]) - set(frame.columns)
        if missing:
            raise LoadError(f"{filename}: отсутствуют столбцы: {', '.join(sorted(missing))}.")
        return frame[SCHEMAS[source]].reset_index(drop=True)
    except LoadError:
        raise
    except Exception as exc:
        # Parser exception details can include uploaded contents: do not echo them.
        raise LoadError(f"{filename}: не удалось прочитать файл; проверьте формат и целостность.") from exc


def load_files(files: dict[str, bytes]) -> dict[str, pd.DataFrame]:
    missing = set(FILENAMES.values()) - set(files)
    if missing:
        raise LoadError(f"Не найдены обязательные файлы: {', '.join(sorted(missing))}.")
    if sum(len(files[name]) for name in FILENAMES.values()) > MAX_TOTAL_BYTES:
        raise LoadError("Общий размер исходных файлов превышает 40 МБ.")
    return {source: read_source(source, files[name]) for source, name in FILENAMES.items()}


def load_zip(data: bytes) -> dict[str, pd.DataFrame]:
    if len(data) > MAX_ZIP_BYTES:
        raise LoadError("Размер ZIP превышает 20 МБ.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = _safe_members(archive, MAX_TOTAL_BYTES)
            files = {}
            for info in members:
                if info.is_dir():
                    continue
                name = PurePosixPath(info.filename).name
                if name in FILENAMES.values():
                    if name in files:
                        raise LoadError(f"В ZIP найдено несколько файлов с именем {name}.")
                    if info.file_size > MAX_FILE_BYTES:
                        raise LoadError(f"{name}: размер превышает 10 МБ.")
                    files[name] = archive.read(info)
            return load_files(files)
    except LoadError:
        raise
    except Exception as exc:
        raise LoadError("Не удалось прочитать ZIP. Проверьте формат и целостность архива.") from exc
